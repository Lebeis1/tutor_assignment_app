"""Struck-through (canceled) classes. Run: python test_canceled.py"""

import os
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Font

from excel_handler import (
    find_canceled_class_ids,
    find_related_sections,
    load_real_choices,
    load_unstaffed_linked_classes,
)


def build_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Choices Round 1"
    ws.append(["Tier", "1st Choice", "2nd Choice", "3rd Choice", "Subject", "Section",
               "Course Description", "Instructor", "Class Status", "Link Group"])
    ws.append(["Tier 1", "Ann Lee", None, None, "MATH 150", "1", "Calc", "Smith", "Active", None])
    ws.append(["Tier 1", "Bo Chan", None, None, "MATH 350", "2", "DiffEq", "Ho", "Active", None])      # canceled
    ws.append(["Tier 2", None, None, None, "COMP 151", "3 LAB", "DS", "Gentry", "Active", "COMP 151 - A"])  # canceled, no names, labelled
    ws.append(["Tier 2", "Cy Dee", None, None, "COMP 151", "2", "DS", "Gentry", "Active", "COMP 151 - A"])
    ws.append(["Tier 2", "Di Fox", None, None, "MATH 201", "1", "Calc2", "Roe", "Active", None])         # only the tutor name is struck
    strike = Font(strike=True)
    for col in range(1, 11):
        ws.cell(row=3, column=col).font = strike     # whole MATH 350 2 row
        ws.cell(row=4, column=col).font = strike     # whole COMP 151 3 LAB row
    ws.cell(row=6, column=2).font = strike           # just "Di Fox" on MATH 201 1
    path = os.path.join(tempfile.mkdtemp(), "t.xlsx")
    wb.save(path)
    return path


def main():
    path = build_workbook()

    assert find_canceled_class_ids(path) == ["MATH 350 2", "COMP 151 3 LAB"]
    print("PASS: rows whose class name is struck through are found")

    ids = [c["class_id"] for c in load_real_choices(path)]
    assert ids == ["MATH 150 1", "COMP 151 2", "MATH 201 1"], ids
    print("PASS: canceled classes are left out of the matching (a struck tutor name alone does not cancel a class)")

    assert load_unstaffed_linked_classes(path) == []
    print("PASS: a canceled class is not listed as an 'also linked' open class")

    assert "MATH 350 2" not in find_related_sections(path)
    print("PASS: canceled classes are left out of the related-section guess")


if __name__ == "__main__":
    main()