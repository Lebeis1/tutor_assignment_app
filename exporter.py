"""Build the downloadable Excel file from the assignments shown on screen."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CHOICE_LABELS = {1: "1st choice", 2: "2nd choice", 3: "3rd choice"}
MAX_ROWS = 2000


def _text(value, limit=500):
    """A safe string for a cell (never None, never absurdly long)."""
    return "" if value is None else str(value)[:limit]


def _put(sheet, row, col, value):
    """Write a value as plain TEXT, so a name or note beginning with "="
    can never be run by Excel as a formula."""
    cell = sheet.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def validate_export_request(data):
    """Check the JSON the page sends. Returns a cleaned dict, or raises ValueError."""
    if not isinstance(data, dict):
        raise ValueError("Nothing to export.")
    assignments = data.get("assignments")
    if not isinstance(assignments, list) or not assignments:
        raise ValueError("There are no assignments to export.")
    if len(assignments) > MAX_ROWS:
        raise ValueError("Too many rows to export.")

    clean_rows = []
    for item in assignments:
        if not isinstance(item, dict) or not item.get("class"):
            raise ValueError("An assignment row is missing its class name.")
        choice = item.get("choice")
        choice2 = item.get("choice2")
        clean_rows.append(
            {
                "class": _text(item["class"]),
                # The course (e.g. "PSY 300") and the section (e.g. "2")
                # travel separately so the Excel file can give each its
                # own column. Missing: fall back to the whole class name.
                "subject": _text(item.get("subject")) or _text(item["class"]),
                "section": _text(item.get("section")),
                "tier": _text(item.get("tier")),
                "tutor": _text(item.get("tutor")),
                "choice": choice if choice in (1, 2, 3) else None,
                "cat": bool(item.get("cat")),
                "tutor2": _text(item.get("tutor2")),
                "choice2": choice2 if choice2 in (1, 2, 3) else None,
                "cat2": bool(item.get("cat2")),
                "locked": [_text(n, 100) for n in (item.get("locked") or [])[:5]],
                "group": _text(item.get("group")),
                "changed": bool(item.get("changed")),
                "notes": [_text(n, 200) for n in (item.get("notes") or [])[:10]],
            }
        )

    tutors = []
    for item in data.get("tutors") or []:
        if isinstance(item, dict) and item.get("name"):
            tutors.append(
                {
                    "name": _text(item["name"]),
                    "load": int(item.get("load") or 0),
                    "max": int(item.get("max") or 0),
                    "classes": [_text(c) for c in (item.get("classes") or [])[:50]],
                }
            )

    tray = [_text(n) for n in (data.get("tray") or [])[:500]]
    canceled = [_text(n) for n in (data.get("canceled") or [])[:500]]
    return {
        "source": _text(data.get("source"), 200),
        "assignments": clean_rows,
        "tutors": tutors,
        "tray": tray,
        "canceled": canceled,
    }


def build_workbook(request):
    """The assignments as an .xlsx file (bytes)."""
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5BEA")
    open_fill = PatternFill("solid", fgColor="FDE7E5")
    changed_fill = PatternFill("solid", fgColor="E8EEFE")

    def style_header(sheet, widths):
        for col, width in enumerate(widths, start=1):
            sheet.cell(row=1, column=col).font = header_font
            sheet.cell(row=1, column=col).fill = header_fill
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "A2"

    wb = Workbook()

    # --- Assignments (one row per class section) -------------------------
    sheet = wb.active
    sheet.title = "Assignments"
    headers = ["Class", "Section", "Tier", "Assigned Tutor", "Tutor's Choice", "CAT",
               "Second Tutor", "Second Tutor's Choice", "Second Tutor CAT",
               "Link Group", "Changed By You", "Notes"]
    for col, name in enumerate(headers, start=1):
        sheet.cell(row=1, column=col, value=name)
    for r, row in enumerate(request["assignments"], start=2):
        choice = CHOICE_LABELS.get(row["choice"], "Not listed" if row["tutor"] else "")
        choice2 = CHOICE_LABELS.get(row["choice2"], "Not listed" if row["tutor2"] else "")
        values = [
            row["subject"], row["section"], row["tier"], row["tutor"] or "OPEN", choice,
            "CAT" if row["cat"] else "",
            row["tutor2"], choice2 if row["tutor2"] else "",
            "CAT" if row["tutor2"] and row["cat2"] else "",
            row["group"],
            "Yes" if row["changed"] else "", "; ".join(row["notes"]),
        ]
        for c, value in enumerate(values, start=1):
            cell = _put(sheet, r, c, value)
            if not row["tutor"]:
                cell.fill = open_fill
            elif row["changed"]:
                cell.fill = changed_fill
    style_header(sheet, [16, 10, 12, 26, 16, 8, 26, 20, 16, 18, 16, 70])
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(request['assignments']) + 1}"

    # --- By Tutor ---------------------------------------------------------
    sheet = wb.create_sheet("By Tutor")
    for col, name in enumerate(["Tutor", "Classes Held", "Max Classes", "Classes"], start=1):
        sheet.cell(row=1, column=col, value=name)
    tutors = sorted(request["tutors"], key=lambda t: t["name"].lower())
    for r, tutor in enumerate(tutors, start=2):
        _put(sheet, r, 1, tutor["name"])
        sheet.cell(row=r, column=2, value=tutor["load"])
        sheet.cell(row=r, column=3, value=tutor["max"])
        _put(sheet, r, 4, ", ".join(tutor["classes"]))
        if tutor["load"] == 0:
            for c in range(1, 5):
                sheet.cell(row=r, column=c).fill = open_fill
        elif tutor["load"] > tutor["max"]:
            sheet.cell(row=r, column=2).font = Font(bold=True, color="B42318")
    style_header(sheet, [28, 14, 14, 90])

    # --- Summary ----------------------------------------------------------
    sheet = wb.create_sheet("Summary")
    rows = request["assignments"]
    filled = sum(1 for r in rows if r["tutor"])
    summary = [
        ("Source file", request["source"]),
        ("Classes", len(rows)),
        ("Classes with a tutor", filled),
        ("Classes with two tutors", sum(1 for r in rows if r["tutor"] and r["tutor2"])),
        ("Locked placements (from the sheet)", sum(len(r["locked"]) for r in rows)),
        ("Classes still open", len(rows) - filled),
        ("Placements the tutor did not list",
         sum(1 for r in rows if r["tutor"] and r["choice"] is None)
         + sum(1 for r in rows if r["tutor2"] and r["choice2"] is None)),
        ("CAT placements", sum(1 for r in rows if r["cat"]) + sum(1 for r in rows if r["tutor2"] and r["cat2"])),
        ("Classes changed by you", sum(1 for r in rows if r["changed"])),
        ("Tutors with no class", sum(1 for t in tutors if t["load"] == 0)),
        ("Names still in the tray (not placed)", ", ".join(request["tray"]) or "None"),
        ("Canceled classes skipped (struck through)", ", ".join(request["canceled"]) or "None"),
    ]
    for r, (label, value) in enumerate(summary, start=1):
        _put(sheet, r, 1, label).font = Font(bold=True)
        _put(sheet, r, 2, value if isinstance(value, int) else _text(value))
        sheet.cell(row=r, column=2).alignment = Alignment(horizontal="left", wrap_text=True)
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 70

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()