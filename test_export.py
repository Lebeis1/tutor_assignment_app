"""Checks for the Excel export. Run: python test_export.py"""

from io import BytesIO

from openpyxl import load_workbook

from app import app


def sample_payload():
    return {
        "source": "test.xlsx",
        "assignments": [
            {"class": "MATH 150 1", "tier": "Tier 1", "tutor": "Ann Lee", "choice": 1, "cat": False,
             "group": "", "changed": False, "notes": []},
            {"class": "COMP 232 1", "tier": "Tier 2", "tutor": "=1+1", "choice": None, "cat": False,
             "group": "COMP 232 - A", "changed": True, "notes": ["=1+1 did not list this class"]},
            {"class": "PSY 300 3", "tier": "CAT Only", "tutor": None, "choice": None, "cat": False,
             "group": "", "changed": False, "notes": []},
        ],
        "tutors": [
            {"name": "Ann Lee", "load": 1, "max": 2, "classes": ["MATH 150 1"]},
            {"name": "Bo Chan", "load": 0, "max": 2, "classes": []},
        ],
        "tray": ["Bo Chan"],
        "canceled": ["MATH 350 2"],
    }


def main():
    client = app.test_client()

    response = client.post("/export", json=sample_payload())
    assert response.status_code == 200
    assert "spreadsheetml" in response.mimetype
    wb = load_workbook(BytesIO(response.data))
    assert wb.sheetnames == ["Assignments", "By Tutor", "Summary"]
    rows = list(wb["Assignments"].iter_rows(values_only=True))
    assert rows[1][:4] == ("MATH 150 1", "Tier 1", "Ann Lee", "1st choice")
    assert rows[3][2] == "OPEN"
    print("PASS: export has the right sheets and rows (open classes marked OPEN)")

    # A name that starts with "=" must be stored as text, never as a formula.
    cell = wb["Assignments"]["C3"]
    assert cell.value == "=1+1" and cell.data_type == "s", (cell.value, cell.data_type)
    print("PASS: a name like =1+1 is stored as plain text, not a formula")

    summary = dict(wb["Summary"].iter_rows(values_only=True))
    assert summary["Classes with a tutor"] == 2 and summary["Classes still open"] == 1
    assert summary["Names still in the tray (not placed)"] == "Bo Chan"
    assert summary["Canceled classes skipped (struck through)"] == "MATH 350 2"
    print("PASS: summary counts, tray names and skipped canceled classes are correct")

    for bad in (None, {}, {"assignments": []}, {"assignments": [{"tier": "x"}]}):
        r = client.post("/export", json=bad)
        assert r.status_code == 400 and "error" in r.get_json(), bad
    print("PASS: empty or malformed requests are rejected with a clear message")


if __name__ == "__main__":
    main()